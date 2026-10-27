#!/usr/bin/env python3
"""29 画风提示语对照表：每组一行，A 效果、B 效果（缩略图加一句话）、提示语、容易失败的点。
用法: python3 prompt_table.py <交付根> <标准片段目录(color_clips.py 640×360 输出)>
输出 交付根/脚本/29画风提示语对照表.md、.xlsx 和 提示语对照表_缩略图/（每段 8 秒窗第 4 秒一帧）。
文字都来自 script_data.py（card、FX），顺序是色序 v4。
"""
import os, sys
import cv2
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import script_data as D

OUT, CL = sys.argv[1:3]
n = len(D.S); TD = os.path.join(OUT, '脚本/提示语对照表_缩略图'); os.makedirs(TD, exist_ok=True)
for f in os.listdir(TD): os.remove(os.path.join(TD, f))
thumbs = {}
for k, s in enumerate(D.S):
    for L in 'AB':
        cap = cv2.VideoCapture(os.path.join(CL, f'{s["id"]}_{L}.mp4')); cap.set(cv2.CAP_PROP_POS_FRAMES, 120); ok, fr = cap.read(); assert ok, s['id']
        p = os.path.join(TD, f'{k + 1:02d}_{s["id"]}_{L}.jpg'); cv2.imwrite(p, cv2.resize(fr, (480, 270)), [cv2.IMWRITE_JPEG_QUALITY, 82]); thumbs[(s['id'], L)] = p

# —— Markdown ——
rel = lambda p: os.path.relpath(p, os.path.join(OUT, '脚本'))
L = [f'# {n} 画风提示语对照表', '',
     f'顺序是色序 v4，从 {D.S[0]["id"]} {D.S[0]["name"]} 开场，和 `{n}画风口播排版_9x16.mp4` 一致。每行一组：A 和 B 是同一画风的两个例子，缩略图取 8 秒窗的第 4 秒。'
     '提示语是屏幕上那张卡片的原话，可以直接把你的主题填进去；完整版（加了画面关键词和输出规格）在 `29画风口播脚本.md`。'
     '容易失败的点合并了来源配方里的注意事项和这轮制作、返修踩过的坑。表格版是同目录的 `.xlsx`，多一列完整提示语。', '',
     '| 序 | 画风 | A 效果 | B 效果 | 提示语 | 容易失败的点 |', '|---|---|---|---|---|---|']
cell = lambda t: t.replace('|', '｜')
for k, s in enumerate(D.S):
    img = lambda Lr: f'<img src="{rel(thumbs[(s["id"], Lr)])}" width="240">'
    L.append(f'| {k + 1:02d} | **{s["name"]}**<br>{s["id"]} | {img("A")}<br>{cell(s["fx_a"])} | {img("B")}<br>{cell(s["fx_b"])} | {cell(s["card"])} | {cell(s["fail"])} |')
md = os.path.join(OUT, f'脚本/{n}画风提示语对照表.md'); open(md, 'w').write('\n'.join(L) + '\n')

# —— Excel（单强调色、1px 边框） ——
wb = Workbook(); ws = wb.active; ws.title = f'{n}画风提示语'
head = ['序', '编号', '画风', 'A 效果图', 'A 效果', 'B 效果图', 'B 效果', '提示语', '容易失败的点', '完整提示语（复制给 Opus 5.5）']
widths = [5, 7, 14, 36, 34, 36, 34, 40, 42, 70]
thin = Side(style='thin', color='111111'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
for j, (h, w) in enumerate(zip(head, widths), 1):
    c = ws.cell(1, j, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='FF5414'); c.border = bd
    c.alignment = Alignment(horizontal='center', vertical='center'); ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 24; ws.freeze_panes = 'D2'
for k, s in enumerate(D.S):
    r = k + 2; ws.row_dimensions[r].height = 112
    vals = [k + 1, s['id'], s['name'], '', s['fx_a'], '', s['fx_b'], s['card'], s['fail'], D.full_prompt(s)]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.border = bd; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(r, 3).font = Font(bold=True); ws.cell(r, 9).font = Font(color='C2410C')
    for col, Lr in (('D', 'A'), ('F', 'B')):
        im = XImage(thumbs[(s['id'], Lr)]); im.width, im.height = 248, 140; ws.add_image(im, f'{col}{r}')
xl = os.path.join(OUT, f'脚本/{n}画风提示语对照表.xlsx'); wb.save(xl)
print(md, xl, len(D.S))
